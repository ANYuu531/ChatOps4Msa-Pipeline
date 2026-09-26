#!/usr/bin/env python3
"""Turns a printed deck PDF into a .pptx of full-bleed page images, for Google Slides.

The HTML deck is the source of truth; this exists only because Google Slides cannot
import a PDF. Each page becomes one 16:9 slide holding that page as a picture, so the
layout is exactly what the PDF shows — nothing reflows, nothing to fix by hand. The
slides are therefore NOT editable as text, which is the accepted trade-off.

Per-page speaker notes are pulled out of the逐頁講稿 markdown (`## P<n> · …` sections)
and put in each slide's notes field, so the script travels with the deck.

    python3 docs/deck-tools/pdf_to_pptx.py <deck.pdf> <out.pptx> [script.md]

Needs `pdftoppm` (poppler) on PATH and python-pptx.
"""
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Emu, Inches

DPI = 150  # 960x540 pt page -> 2000x1125 px: crisp on a projector, still a sane file size


def page_images(pdf: pathlib.Path, workdir: pathlib.Path) -> list[pathlib.Path]:
    if not shutil.which("pdftoppm"):
        sys.exit("pdftoppm not found — install poppler (brew install poppler)")
    subprocess.run(["pdftoppm", "-png", "-r", str(DPI), str(pdf), str(workdir / "page")], check=True)
    return sorted(workdir.glob("page-*.png"))


def notes_by_page(script: pathlib.Path | None) -> dict[int, str]:
    """`## P<n> · title` … up to the next `## ` — the spoken script for that page."""
    if script is None or not script.exists():
        return {}
    out: dict[int, str] = {}
    for block in re.split(r"\n## ", script.read_text()):
        m = re.match(r"P(\d+)\s*·\s*(.*)", block)
        if not m:
            continue
        body = block.split("\n", 1)[1] if "\n" in block else ""
        body = body.replace("**", "").replace("`", "").strip()
        body = re.sub(r"\n---\s*$", "", body).strip()
        out[int(m.group(1))] = f"P{m.group(1)} · {m.group(2)}\n\n{body}"
    return out


def main(pdf_path: str, out_path: str, script_path: str | None) -> None:
    pdf = pathlib.Path(pdf_path)
    notes = notes_by_page(pathlib.Path(script_path) if script_path else None)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    blank = prs.slide_layouts[6]

    with tempfile.TemporaryDirectory() as tmp:
        images = page_images(pdf, pathlib.Path(tmp))
        if not images:
            sys.exit("no pages rendered")
        for i, image in enumerate(images, start=1):
            slide = prs.slides.add_slide(blank)
            slide.shapes.add_picture(str(image), Emu(0), Emu(0),
                                     width=prs.slide_width, height=prs.slide_height)
            if i in notes:
                slide.notes_slide.notes_text_frame.text = notes[i]

    prs.save(out_path)
    print(f"wrote {out_path} — {len(images)} slides, {len(notes)} with speaker notes")


if __name__ == "__main__":
    if not 3 <= len(sys.argv) <= 4:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2], sys.argv[3] if len(sys.argv) > 3 else None)
